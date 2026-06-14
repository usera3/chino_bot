#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Local bridge for exposing selected chino_bot capabilities to OpenClaw."""

from __future__ import annotations

import argparse
import asyncio
import contextlib
import email
from email.header import decode_header, Header
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email import encoders
import imaplib
import json
import os
from pathlib import Path
import secrets
import re
import shutil
import smtplib
import subprocess
import ssl
import sys
import time
from typing import Any, Callable
import urllib.error
import urllib.parse
import urllib.request


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ENV_PATH = PROJECT_ROOT / ".env"
DOWNLOADS_DIR = Path.home() / "Downloads"
OPENCLAW_HOME = Path.home() / ".openclaw"
OPENCLAW_CONFIG_PATH = OPENCLAW_HOME / "openclaw.json"
OPENCLAW_QQ_MEDIA_PROXY_DIR = OPENCLAW_HOME / "canvas" / "qq-media"
OPENCLAW_QQ_MEDIA_PROXY_PATH = "/__openclaw__/canvas/qq-media"
TEMP_IMAGE_DIR = PROJECT_ROOT / "data" / "temp_images"
TAVILY_API_URL = "https://api.tavily.com/search"
RENDER_BROWSER_CDP_HOST = "127.0.0.1"
RENDER_BROWSER_CDP_PORT = int(os.getenv("OPENCLAW_RENDER_BROWSER_CDP_PORT", "39223"))
RENDER_BROWSER_PROFILE_DIR = PROJECT_ROOT / "data" / "render_browser_profile"


def _strip_matching_quotes(value: str) -> str:
    if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
        return value[1:-1]
    return value


def load_project_env_file(path: Path) -> tuple[dict[str, str], set[str]]:
    loaded: dict[str, str] = {}
    applied: set[str] = set()
    try:
        if not path.exists():
            return loaded, applied
        for raw_line in path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue
            if line.startswith("export "):
                line = line[7:].strip()
            if "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            if not key:
                continue
            value = _strip_matching_quotes(value.strip())
            loaded[key] = value
            if not os.getenv(key, "").strip():
                os.environ[key] = value
                applied.add(key)
        for upper_name in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "NO_PROXY"):
            value = os.getenv(upper_name, "").strip()
            lower_name = upper_name.lower()
            if value and not os.getenv(lower_name, "").strip():
                os.environ[lower_name] = value
    except Exception:
        return loaded, applied
    return loaded, applied


PROJECT_ENV, PROJECT_ENV_APPLIED = load_project_env_file(PROJECT_ENV_PATH)


def ok(result: Any) -> dict[str, Any]:
    return {"ok": True, "result": result}


def fail(message: str, *, details: Any = None) -> dict[str, Any]:
    payload: dict[str, Any] = {"ok": False, "error": message}
    if details is not None:
        payload["details"] = details
    return payload


def decode_mime_header(value: str) -> str:
    if not value:
        return ""
    decoded = []
    for part, encoding in decode_header(value):
        if isinstance(part, bytes):
            candidates = [encoding, "utf-8", "gbk", "gb2312", "latin1"]
            text = None
            for candidate in candidates:
                if not candidate:
                    continue
                try:
                    text = part.decode(candidate, errors="ignore")
                    break
                except Exception:
                    continue
            decoded.append(text if text is not None else part.decode("utf-8", errors="ignore"))
        else:
            decoded.append(part)
    return "".join(decoded).strip()


def load_openclaw_config() -> dict[str, Any]:
    try:
        if not OPENCLAW_CONFIG_PATH.exists():
            return {}
        payload = json.loads(OPENCLAW_CONFIG_PATH.read_text(encoding="utf-8"))
        return payload if isinstance(payload, dict) else {}
    except Exception:
        return {}


def is_executable(path: Path) -> bool:
    try:
        return path.is_file() and os.access(path, os.X_OK)
    except Exception:
        return False


def resolve_openclaw_gateway_port() -> int:
    config = load_openclaw_config()
    gateway = config.get("gateway")
    if isinstance(gateway, dict):
        port = gateway.get("port")
        if isinstance(port, int) and port > 0:
            return port
    return 18789


def resolve_browser_executable_path() -> Path | None:
    config = load_openclaw_config()
    config_browser = config.get("browser")
    candidates: list[Path] = []

    if isinstance(config_browser, dict):
        configured = str(config_browser.get("executablePath", "")).strip()
        if configured:
            candidates.append(Path(configured).expanduser())

    for env_name in (
        "OPENCLAW_BROWSER_EXECUTABLE_PATH",
        "BROWSER_EXECUTABLE_PATH",
        "PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH",
    ):
        value = os.getenv(env_name, "").strip()
        if value:
            candidates.append(Path(value).expanduser())

    candidates.extend(
        [
            Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"),
            Path("/Applications/Chromium.app/Contents/MacOS/Chromium"),
            Path("/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"),
            Path("/Applications/Brave Browser.app/Contents/MacOS/Brave Browser"),
        ]
    )

    playwright_cache = Path.home() / "Library" / "Caches" / "ms-playwright"
    if playwright_cache.exists():
        for chromium_dir in sorted(playwright_cache.glob("chromium-*"), reverse=True):
            candidates.extend(
                [
                    chromium_dir
                    / "chrome-mac-arm64"
                    / "Google Chrome for Testing.app"
                    / "Contents"
                    / "MacOS"
                    / "Google Chrome for Testing",
                    chromium_dir / "chrome-mac" / "Chromium",
                    chromium_dir / "chrome-linux" / "chrome",
                    chromium_dir / "chrome-win" / "chrome.exe",
                    chromium_dir / "chrome-win64" / "chrome.exe",
                    chromium_dir / "headless_shell",
                    chromium_dir / "chrome-mac-arm64" / "headless_shell",
                ]
            )

    seen: set[str] = set()
    for candidate in candidates:
        normalized = str(candidate)
        if normalized in seen:
            continue
        seen.add(normalized)
        if is_executable(candidate):
            return candidate
    return None


def publish_openclaw_media_proxy(source_file_path: Path, prefix: str) -> str:
    if not source_file_path.exists():
        raise FileNotFoundError(f"render output not found: {source_file_path}")
    ext = source_file_path.suffix or ".png"
    output_name = f"{prefix}-{int(time.time() * 1000)}-{os.getpid()}{ext}"
    target_path = OPENCLAW_QQ_MEDIA_PROXY_DIR / output_name
    OPENCLAW_QQ_MEDIA_PROXY_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_file_path, target_path)
    port = resolve_openclaw_gateway_port()
    return f"http://127.0.0.1:{port}{OPENCLAW_QQ_MEDIA_PROXY_PATH}/{output_name}"


def resolve_input_path(raw_path: Any) -> Path:
    candidate = Path(str(raw_path or "").strip()).expanduser()
    if candidate.is_absolute():
        return candidate
    direct = candidate.resolve()
    if direct.exists():
        return direct
    project_candidate = (PROJECT_ROOT / candidate).resolve()
    if project_candidate.exists():
        return project_candidate
    downloads_candidate = (DOWNLOADS_DIR / candidate.name).resolve()
    if downloads_candidate.exists():
        return downloads_candidate
    return direct


def resolve_downloads_output_path(raw_path: Any, default_name: str) -> Path:
    candidate = str(raw_path or "").strip()
    filename = Path(candidate).name if candidate else default_name
    return DOWNLOADS_DIR / filename


def ensure_temp_image_path(prefix: str, suffix: str = ".png") -> Path:
    TEMP_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    unique = f"{time.time_ns()}_{os.getpid()}_{secrets.token_hex(4)}"
    return TEMP_IMAGE_DIR / f"{prefix}_{unique}{suffix}"


@contextlib.contextmanager
def bypass_proxy_env():
    keys = [
        "HTTP_PROXY",
        "HTTPS_PROXY",
        "ALL_PROXY",
        "NO_PROXY",
        "http_proxy",
        "https_proxy",
        "all_proxy",
        "no_proxy",
    ]
    previous = {key: os.environ.get(key) for key in keys}
    try:
        for key in keys:
            os.environ.pop(key, None)
        yield
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


def render_browser_cdp_base_url() -> str:
    return f"http://{RENDER_BROWSER_CDP_HOST}:{RENDER_BROWSER_CDP_PORT}"


def is_render_browser_ready(timeout: float = 0.35) -> bool:
    try:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(
            f"{render_browser_cdp_base_url()}/json/version",
            timeout=timeout,
        ) as response:
            return response.status == 200
    except Exception:
        return False


def resolve_render_browser_cdp_ws_url(timeout: float = 0.5) -> str | None:
    try:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(
            f"{render_browser_cdp_base_url()}/json/version",
            timeout=timeout,
        ) as response:
            payload = json.loads(response.read().decode("utf-8", errors="ignore"))
        ws_url = str(payload.get("webSocketDebuggerUrl", "")).strip()
        return ws_url or None
    except Exception:
        return None


def launch_render_browser_process(browser_executable_path: Path) -> None:
    launch_args = [
        str(browser_executable_path),
        "--headless",
        f"--remote-debugging-port={RENDER_BROWSER_CDP_PORT}",
        "--disable-dev-shm-usage",
        "--disable-background-networking",
        "--disable-default-apps",
        "--disable-extensions",
        "--disable-features=Translate,OptimizationHints,MediaRouter",
        "--no-sandbox",
        "--disable-setuid-sandbox",
        "--no-first-run",
        "--no-default-browser-check",
        "about:blank",
    ]
    subprocess.Popen(
        launch_args,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        stdin=subprocess.DEVNULL,
        start_new_session=True,
        close_fds=True,
    )


def ensure_render_browser_ready(browser_executable_path: Path) -> bool:
    if is_render_browser_ready():
        return True
    launch_render_browser_process(browser_executable_path)
    deadline = time.time() + 6.0
    while time.time() < deadline:
        if is_render_browser_ready(timeout=0.5):
            return True
        time.sleep(0.1)
    return False


def resolve_playwright_proxy_config() -> dict[str, str] | None:
    use_proxy = os.getenv("USE_PROXY", "false").strip().lower() == "true"
    if not use_proxy:
        return None
    http_proxy = os.getenv("HTTP_PROXY", "").strip()
    if not http_proxy:
        return None
    return {"server": http_proxy}


def is_retryable_playwright_proxy_error(error: Exception) -> bool:
    message = str(error or "")
    if not message:
        return False
    retry_markers = (
        "ERR_PROXY_CONNECTION_FAILED",
        "ERR_TUNNEL_CONNECTION_FAILED",
        "ERR_NO_SUPPORTED_PROXIES",
        "proxy connection failed",
        "proxyconnect tcp",
    )
    return any(marker in message for marker in retry_markers)


async def _wait_for_page_render_ready(
    page: Any,
    *,
    wait_ms: int,
    require_network_idle: bool,
    network_idle_timeout_ms: int = 5_000,
) -> None:
    if require_network_idle:
        try:
            await page.wait_for_load_state("networkidle", timeout=network_idle_timeout_ms)
        except Exception:
            pass
    await page.evaluate(
        """async () => {
            if (document.fonts && document.fonts.ready) {
                await document.fonts.ready;
            }
            const images = Array.from(document.images || []);
            await Promise.allSettled(images.map((img) => {
                if (img.complete) {
                    return Promise.resolve();
                }
                return new Promise((resolve) => {
                    img.addEventListener("load", resolve, { once: true });
                    img.addEventListener("error", resolve, { once: true });
                });
            }));
            await Promise.allSettled(images.map((img) => {
                if (typeof img.decode === "function") {
                    return img.decode().catch(() => {});
                }
                return Promise.resolve();
            }));
            await new Promise((resolve) => requestAnimationFrame(() => resolve(null)));
            await new Promise((resolve) => requestAnimationFrame(() => resolve(null)));
        }"""
    )
    if wait_ms > 0:
        await page.wait_for_timeout(wait_ms)


async def _wait_for_cdp_default_context(browser: Any, timeout_ms: int = 1200) -> Any | None:
    deadline = time.monotonic() + max(timeout_ms, 0) / 1000.0
    while time.monotonic() < deadline:
        contexts = list(browser.contexts)
        if contexts:
            return contexts[0]
        await asyncio.sleep(0.05)
    return None


async def _acquire_reused_page(context: Any) -> tuple[Any, bool]:
    try:
        return await context.new_page(), True
    except Exception:
        existing_pages = list(context.pages)
        if existing_pages:
            return existing_pages[0], False
        raise


async def _open_render_page(
    *,
    playwright: Any,
    browser_executable_path: Path,
    width: int,
    height: int,
    device_scale_factor: float,
    proxy_config: dict[str, str] | None,
    prefer_reuse: bool,
) -> tuple[Any, Any, Any, bool, bool, bool]:
    reused = False
    browser = None
    context = None
    created_context = False

    if prefer_reuse and proxy_config is None and ensure_render_browser_ready(browser_executable_path):
        last_cdp_error: Exception | None = None
        for _ in range(8):
            try:
                cdp_ws_url = resolve_render_browser_cdp_ws_url()
                if not cdp_ws_url:
                    raise RuntimeError("cdp browser is alive but websocket debugger url is unavailable")
                browser = await playwright.chromium.connect_over_cdp(cdp_ws_url)
                context = (
                    browser.contexts[0]
                    if browser.contexts
                    else await _wait_for_cdp_default_context(browser)
                )
                if context is None:
                    raise RuntimeError("cdp browser connected but no default context became available")
                reused = True
                break
            except Exception as exc:
                last_cdp_error = exc
                if browser is not None:
                    try:
                        await browser.close()
                    except Exception:
                        pass
                browser = None
                context = None
                reused = False
                created_context = False
                await asyncio.sleep(0.1)
        if browser is None and last_cdp_error is not None:
            pass

    if browser is None or context is None:
        browser = await playwright.chromium.launch(
            executable_path=str(browser_executable_path),
            headless=True,
            args=["--disable-dev-shm-usage", "--no-sandbox", "--disable-setuid-sandbox"],
            proxy=proxy_config,
        )
        context = await browser.new_context(
            viewport={"width": width, "height": height},
            device_scale_factor=device_scale_factor,
            user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36",
        )
        created_context = True

    if reused:
        page, owns_page = await _acquire_reused_page(context)
    else:
        page = await context.new_page()
        owns_page = True
    return browser, context, page, reused, created_context, owns_page


def should_require_network_idle_for_url(url: str) -> bool:
    try:
        parsed = urllib.parse.urlparse(url)
    except Exception:
        return True
    if parsed.scheme == "file":
        return False
    host = (parsed.hostname or "").strip().lower()
    if host in {"127.0.0.1", "localhost", "::1", "0.0.0.0"}:
        return False
    return True


def should_use_proxy_for_url(url: str) -> bool:
    try:
        parsed = urllib.parse.urlparse(url)
    except Exception:
        return True
    if parsed.scheme == "file":
        return False
    host = (parsed.hostname or "").strip().lower()
    if host in {"127.0.0.1", "localhost", "::1", "0.0.0.0"}:
        return False
    return True


def maybe_limit_text_output(text: str, max_chars: int = 60_000) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n…(内容过长，已截断)"


def coerce_bool(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"1", "true", "yes", "on"}:
            return True
        if normalized in {"0", "false", "no", "off"}:
            return False
    return default


def coerce_int(
    value: Any,
    *,
    default: int,
    minimum: int | None = None,
    maximum: int | None = None,
) -> int:
    try:
        parsed = int(value)
    except Exception:
        parsed = default
    if minimum is not None and parsed < minimum:
        parsed = minimum
    if maximum is not None and parsed > maximum:
        parsed = maximum
    return parsed


def normalize_string_list(value: Any, *, max_items: int = 10) -> list[str]:
    if not isinstance(value, list):
        return []
    normalized: list[str] = []
    for item in value:
        text = str(item).strip()
        if not text:
            continue
        normalized.append(text)
        if len(normalized) >= max_items:
            break
    return normalized


def mask_url_for_debug(raw_url: str) -> str | None:
    value = raw_url.strip()
    if not value:
        return None
    try:
        parsed = urllib.parse.urlsplit(value)
        if not parsed.scheme or not parsed.hostname:
            return value
        port = f":{parsed.port}" if parsed.port else ""
        return f"{parsed.scheme}://{parsed.hostname}{port}"
    except Exception:
        return value


def resolve_env_source(env_name: str) -> str:
    if env_name in PROJECT_ENV_APPLIED:
        return "project_env"
    if os.getenv(env_name, "").strip():
        return "process_env"
    return "missing"


def resolve_tavily_api_key() -> tuple[str, str]:
    value = os.getenv("TAVILY_API_KEY", "").strip()
    if not value:
        return "", "missing"
    return value, resolve_env_source("TAVILY_API_KEY")


def resolve_project_proxy_settings() -> dict[str, str]:
    proxies: dict[str, str] = {}
    for scheme, env_names in (
        ("http", ("HTTP_PROXY", "http_proxy")),
        ("https", ("HTTPS_PROXY", "https_proxy")),
    ):
        for env_name in env_names:
            value = os.getenv(env_name, "").strip()
            if value:
                proxies[scheme] = value
                break
    return proxies


def build_tavily_opener(*, use_proxy: bool) -> tuple[urllib.request.OpenerDirector, dict[str, Any]]:
    handlers: list[Any] = []
    proxies = resolve_project_proxy_settings() if use_proxy else {}
    if use_proxy:
        if not proxies:
            raise RuntimeError("project proxy requested but HTTP_PROXY / HTTPS_PROXY is not configured")
    handlers.append(urllib.request.ProxyHandler(proxies if use_proxy else {}))

    disable_ssl_verify = coerce_bool(
        os.getenv("TAVILY_DISABLE_SSL_VERIFY"),
        default=True,
    )
    ssl_context = ssl.create_default_context()
    if disable_ssl_verify:
        ssl_context.check_hostname = False
        ssl_context.verify_mode = ssl.CERT_NONE
    handlers.append(urllib.request.HTTPSHandler(context=ssl_context))

    return urllib.request.build_opener(*handlers), {
        "mode": "project_proxy" if use_proxy else "direct",
        "proxy_enabled": use_proxy,
        "http_proxy": mask_url_for_debug(proxies.get("http", "")),
        "https_proxy": mask_url_for_debug(proxies.get("https", "")),
        "ssl_verify": not disable_ssl_verify,
    }


def perform_tavily_request(
    *,
    request_payload: dict[str, Any],
    timeout_seconds: int,
    use_proxy: bool,
) -> tuple[dict[str, Any], dict[str, Any]]:
    opener, transport_meta = build_tavily_opener(use_proxy=use_proxy)
    request = urllib.request.Request(
        TAVILY_API_URL,
        data=json.dumps(request_payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "Accept": "application/json",
            "Content-Type": "application/json",
            "User-Agent": "openclaw-chinobot-bridge/2026.04.07",
        },
        method="POST",
    )

    try:
        with opener.open(request, timeout=timeout_seconds) as response:
            body = response.read().decode("utf-8", errors="replace")
            payload = json.loads(body)
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        snippet = maybe_limit_text_output(body.strip(), max_chars=1_200)
        raise RuntimeError(f"Tavily HTTP {exc.code}: {snippet}") from exc
    except urllib.error.URLError as exc:
        reason = exc.reason if isinstance(exc.reason, str) else repr(exc.reason)
        raise RuntimeError(f"Tavily request failed: {reason}") from exc
    except TimeoutError as exc:
        raise RuntimeError("Tavily request timed out") from exc
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Tavily returned invalid JSON: {exc}") from exc

    if not isinstance(payload, dict):
        raise RuntimeError("Tavily returned a non-object response")
    if payload.get("error"):
        raise RuntimeError(str(payload.get("error")))
    return payload, transport_meta


def register_reportlab_font() -> tuple[str, str | None]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidate_fonts = [
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Hiragino Sans GB.ttc"),
        Path("/System/Library/Fonts/STHeiti Medium.ttc"),
        Path("/System/Library/Fonts/Supplemental/Songti.ttc"),
    ]

    for font_path in candidate_fonts:
        if not font_path.exists():
            continue
        try:
            pdfmetrics.registerFont(TTFont("BridgeCJK", str(font_path)))
            return "BridgeCJK", str(font_path)
        except Exception:
            continue
    return "Helvetica", None


async def _render_html_to_png(
    *,
    html_code: str,
    output_path: Path,
    width: int,
    height: int,
    device_scale_factor: float,
    wait_ms: int,
    full_page: bool,
    browser_executable_path: Path,
) -> str:
    from playwright.async_api import async_playwright

    requires_network_idle = bool(re.search(r"""https?://|src=|href=|url\(""", html_code, re.IGNORECASE))

    async with async_playwright() as playwright:
        browser, context, page, reused, created_context, owns_page = await _open_render_page(
            playwright=playwright,
            browser_executable_path=browser_executable_path,
            width=width,
            height=height,
            device_scale_factor=device_scale_factor,
            proxy_config=None,
            prefer_reuse=True,
        )
        try:
            await page.set_content(html_code, wait_until="domcontentloaded")
            await _wait_for_page_render_ready(
                page,
                wait_ms=wait_ms,
                require_network_idle=requires_network_idle,
            )
            await page.screenshot(
                path=str(output_path),
                type="png",
                full_page=full_page,
                animations="disabled",
            )
        finally:
            if reused and not owns_page:
                try:
                    await page.goto("about:blank", wait_until="domcontentloaded", timeout=2_000)
                except Exception:
                    pass
            else:
                await page.close()
            if created_context:
                await context.close()
            if not reused:
                await browser.close()
    return "cdp-reuse" if reused else "ephemeral-launch"


async def _capture_web_screenshot(
    *,
    url: str,
    output_path: Path,
    width: int,
    height: int,
    full_page: bool,
    wait_ms: int,
    browser_executable_path: Path,
    proxy_config: dict[str, str] | None,
) -> str:
    from playwright.async_api import async_playwright
    require_network_idle = should_require_network_idle_for_url(url)

    async with async_playwright() as playwright:
        browser, context, page, reused, created_context, owns_page = await _open_render_page(
            playwright=playwright,
            browser_executable_path=browser_executable_path,
            width=width,
            height=height,
            device_scale_factor=1.0,
            proxy_config=proxy_config,
            prefer_reuse=proxy_config is None,
        )
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=30_000)
            await _wait_for_page_render_ready(
                page,
                wait_ms=wait_ms,
                require_network_idle=require_network_idle,
                network_idle_timeout_ms=1_200,
            )
            await page.screenshot(
                path=str(output_path),
                type="png",
                full_page=full_page,
                animations="disabled",
            )
        finally:
            if reused and not owns_page:
                try:
                    await page.goto("about:blank", wait_until="domcontentloaded", timeout=2_000)
                except Exception:
                    pass
            else:
                await page.close()
            if created_context:
                await context.close()
            if not reused:
                await browser.close()
    return "cdp-reuse" if reused else "ephemeral-launch"


def tool_render_html(params: dict[str, Any]) -> dict[str, Any]:
    html_code = str(params.get("html_code", "")).strip()
    if not html_code:
        raise ValueError("html_code 不能为空")

    try:
        width = int(params.get("width", 800) or 800)
        height = int(params.get("height", 600) or 600)
        wait_ms = int(params.get("wait_ms", 150) or 150)
    except Exception as exc:
        raise ValueError(f"width / height / wait_ms 必须是数字: {exc}") from exc

    width = max(64, min(width, 4096))
    height = max(64, min(height, 4096))
    wait_ms = max(0, min(wait_ms, 10_000))

    try:
        device_scale_factor = float(params.get("device_scale_factor", 2) or 2)
    except Exception as exc:
        raise ValueError(f"device_scale_factor 必须是数字: {exc}") from exc
    device_scale_factor = max(1.0, min(device_scale_factor, 4.0))
    full_page = bool(params.get("full_page", True))

    browser_executable_path = resolve_browser_executable_path()
    if browser_executable_path is None:
        raise RuntimeError(
            "未找到 Chromium-compatible browser。可在 ~/.openclaw/openclaw.json 的 browser.executablePath 配置浏览器，"
            "或设置 OPENCLAW_BROWSER_EXECUTABLE_PATH / PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH。"
        )

    output_path = ensure_temp_image_path("render")

    try:
        browser_mode = asyncio.run(
            _render_html_to_png(
                html_code=html_code,
                output_path=output_path,
                width=width,
                height=height,
                device_scale_factor=device_scale_factor,
                wait_ms=wait_ms,
                full_page=full_page,
                browser_executable_path=browser_executable_path,
            )
        )
    except Exception as exc:
        raise RuntimeError(f"HTML 渲染失败：{exc}") from exc

    file_size = output_path.stat().st_size if output_path.exists() else 0
    media_url = publish_openclaw_media_proxy(output_path, "chinobot-render")
    return {
        "ok": True,
        "message": "✅ HTML 渲染成功",
        "file_path": str(output_path),
        "path": str(output_path),
        "media_url": media_url,
        "mediaUrl": media_url,
        "primary_media_url": media_url,
        "primaryMediaUrl": media_url,
        "width": width,
        "height": height,
        "wait_ms": wait_ms,
        "device_scale_factor": device_scale_factor,
        "full_page": full_page,
        "file_size_bytes": file_size,
        "browser_executable_path": str(browser_executable_path),
        "browser_mode": browser_mode,
        "note": "图片会由系统自动发送，请不要重复输出链接、文件路径或 [[media:...]]。",
    }


def tool_read_pdf(params: dict[str, Any]) -> str:
    try:
        from PyPDF2 import PdfReader
    except Exception as exc:
        return f"❌ 读取 PDF 失败：缺少 PyPDF2 依赖（{exc}）"

    file_path = resolve_input_path(params.get("file_path"))
    if not file_path.exists():
        return f"❌ 文件不存在: {file_path}"

    max_pages_raw = params.get("max_pages")
    max_pages = None
    if max_pages_raw not in (None, ""):
        try:
            max_pages = max(1, int(max_pages_raw))
        except Exception as exc:
            return f"❌ max_pages 参数无效：{exc}"

    try:
        reader = PdfReader(str(file_path))
        total_pages = len(reader.pages)
        pages_to_read = min(max_pages, total_pages) if max_pages else total_pages
        text_content: list[str] = []
        for index in range(pages_to_read):
            page_text = (reader.pages[index].extract_text() or "").strip()
            if page_text:
                text_content.append(f"[第 {index + 1} 页]\n{page_text}")
        if not text_content:
            return "❌ 无法从 PDF 中提取文本（可能是扫描版或图片 PDF）"
        summary = (
            "📄 PDF 文档信息\n"
            f"📁 文件：{file_path.name}\n"
            f"📊 总页数：{total_pages}\n"
            f"📖 已读取：{pages_to_read} 页\n"
            f"{'=' * 50}\n\n"
        )
        return maybe_limit_text_output(summary + "\n\n".join(text_content))
    except Exception as exc:
        return f"❌ 读取 PDF 失败：{exc}"


def tool_read_excel(params: dict[str, Any]) -> str:
    try:
        import pandas as pd
    except Exception as exc:
        return f"❌ 读取 Excel 失败：缺少 pandas/openpyxl 依赖（{exc}）"

    file_path = resolve_input_path(params.get("file_path"))
    if not file_path.exists():
        return f"❌ 文件不存在: {file_path}"

    sheet_name = str(params.get("sheet_name", "")).strip() or None
    max_rows_raw = params.get("max_rows", 100)
    try:
        max_rows = max(1, int(max_rows_raw)) if max_rows_raw not in (None, "") else None
    except Exception as exc:
        return f"❌ max_rows 参数无效：{exc}"

    try:
        read_kwargs: dict[str, Any] = {}
        if sheet_name:
            read_kwargs["sheet_name"] = sheet_name
        if max_rows:
            read_kwargs["nrows"] = max_rows
        df = pd.read_excel(str(file_path), **read_kwargs)
        workbook = pd.ExcelFile(str(file_path))
        current_sheet = sheet_name or workbook.sheet_names[0]
        summary = (
            "📊 Excel 表格信息\n"
            f"📁 文件：{file_path.name}\n"
            f"📋 工作表：{', '.join(workbook.sheet_names)}\n"
            f"📏 当前表：{current_sheet}\n"
            f"📊 行数：{len(df)} / 列数：{len(df.columns)}\n"
            f"{'=' * 50}\n\n"
        )
        table_str = df.to_string(index=False, max_rows=max_rows)
        return maybe_limit_text_output(summary + table_str)
    except Exception as exc:
        return f"❌ 读取 Excel 失败：{exc}"


def tool_create_excel(params: dict[str, Any]) -> str:
    try:
        import pandas as pd
    except Exception as exc:
        return f"❌ 创建 Excel 失败：缺少 pandas/openpyxl 依赖（{exc}）"

    file_path = resolve_downloads_output_path(params.get("file_path"), "chinobot-output.xlsx")
    data_value = params.get("data")
    if data_value in (None, ""):
        return "❌ data 不能为空"
    sheet_name = str(params.get("sheet_name", "Sheet1") or "Sheet1").strip() or "Sheet1"

    try:
        data_list = json.loads(data_value) if isinstance(data_value, str) else data_value
    except json.JSONDecodeError as exc:
        return f"❌ JSON 解析失败：{exc}"

    if not isinstance(data_list, list):
        return "❌ 数据格式错误：必须是列表格式"

    try:
        df = pd.DataFrame(data_list)
        df.to_excel(str(file_path), sheet_name=sheet_name, index=False)
        file_size_kb = file_path.stat().st_size / 1024
        return (
            "✅ Excel 表格创建成功！\n"
            f"📁 文件：{file_path}\n"
            f"📊 大小：{file_size_kb:.2f} KB\n"
            f"📏 行数：{len(df)} / 列数：{len(df.columns)}"
        )
    except Exception as exc:
        return f"❌ 创建 Excel 失败：{exc}"


def tool_convert_word_to_pdf(params: dict[str, Any]) -> str:
    try:
        from docx import Document
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except Exception as exc:
        return f"❌ Word 转 PDF 失败：缺少依赖（{exc}）"

    word_path = resolve_input_path(params.get("word_path"))
    if not word_path.exists():
        return f"❌ Word 文件不存在: {word_path}"

    pdf_path = resolve_downloads_output_path(
        params.get("pdf_path"),
        f"{word_path.stem}.pdf",
    )

    try:
        font_name, font_path = register_reportlab_font()
        doc = Document(str(word_path))
        story: list[Any] = []
        styles = getSampleStyleSheet()
        normal_style = ParagraphStyle(
            "BridgeChineseNormal",
            parent=styles["Normal"],
            fontName=font_name,
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
        )
        heading_style = ParagraphStyle(
            "BridgeChineseHeading",
            parent=styles["Heading1"],
            fontName=font_name,
            fontSize=18,
            leading=24,
            alignment=TA_CENTER,
            spaceAfter=12,
        )
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = heading_style if para.style.name.startswith("Heading") else normal_style
            story.append(Paragraph(text.replace("\n", "<br/>"), style))
            story.append(Spacer(1, 0.2 * inch))
        if not story:
            return "❌ Word 文档为空，无法转换"
        pdf = SimpleDocTemplate(str(pdf_path), pagesize=A4)
        pdf.build(story)
        if not pdf_path.exists():
            return "❌ PDF 转换失败：文件未生成"
        file_size_kb = pdf_path.stat().st_size / 1024
        font_note = f"\n🔤 字体：{font_path}" if font_path else "\n🔤 字体：Helvetica（可能不完整支持中文）"
        return (
            "✅ Word 转 PDF 成功！\n"
            f"📁 原文件：{word_path.name}\n"
            f"📄 PDF 文件：{pdf_path}\n"
            f"📊 大小：{file_size_kb:.2f} KB"
            f"{font_note}"
        )
    except Exception as exc:
        return f"❌ Word 转 PDF 失败：{exc}"


def tool_convert_pdf_to_word(params: dict[str, Any]) -> str:
    try:
        from pdf2docx import Converter
    except Exception as exc:
        return f"❌ PDF 转 Word 失败：缺少 pdf2docx 依赖（{exc}）"

    pdf_path = resolve_input_path(params.get("pdf_path"))
    if not pdf_path.exists():
        return f"❌ PDF 文件不存在: {pdf_path}"

    word_path = resolve_downloads_output_path(
        params.get("word_path"),
        f"{pdf_path.stem}.docx",
    )

    try:
        converter = Converter(str(pdf_path))
        try:
            converter.convert(str(word_path))
        finally:
            converter.close()
        if not word_path.exists():
            return "❌ Word 转换失败：文件未生成"
        file_size_kb = word_path.stat().st_size / 1024
        return (
            "✅ PDF 转 Word 成功！\n"
            f"📁 原文件：{pdf_path.name}\n"
            f"📝 Word 文件：{word_path}\n"
            f"📊 大小：{file_size_kb:.2f} KB\n"
            "✨ 已转换为可编辑格式"
        )
    except Exception as exc:
        return f"❌ PDF 转 Word 失败：{exc}"


def tool_web_screenshot(params: dict[str, Any]) -> dict[str, Any]:
    url = str(params.get("url", "")).strip()
    if not url:
        raise ValueError("url 不能为空")
    if not url.startswith(("http://", "https://")):
        raise ValueError("URL 格式错误：必须以 http:// 或 https:// 开头")

    try:
        width = int(params.get("width", 1920) or 1920)
        height = int(params.get("height", 1080) or 1080)
        wait_ms = int(params.get("wait_ms", 350) or 350)
    except Exception as exc:
        raise ValueError(f"width / height / wait_ms 必须是数字: {exc}") from exc

    width = max(320, min(width, 4096))
    height = max(240, min(height, 4096))
    wait_ms = max(0, min(wait_ms, 10_000))
    full_page = bool(params.get("full_page", True))

    browser_executable_path = resolve_browser_executable_path()
    if browser_executable_path is None:
        raise RuntimeError(
            "未找到 Chromium-compatible browser。可在 ~/.openclaw/openclaw.json 的 browser.executablePath 配置浏览器，"
            "或设置 OPENCLAW_BROWSER_EXECUTABLE_PATH / PLAYWRIGHT_CHROMIUM_EXECUTABLE_PATH。"
        )

    output_path = ensure_temp_image_path("screenshot")
    proxy_config = resolve_playwright_proxy_config() if should_use_proxy_for_url(url) else None
    proxy_retry_used = False
    proxy_error_message = None

    try:
        browser_mode = asyncio.run(
            _capture_web_screenshot(
                url=url,
                output_path=output_path,
                width=width,
                height=height,
                full_page=full_page,
                wait_ms=wait_ms,
                browser_executable_path=browser_executable_path,
                proxy_config=proxy_config,
            )
        )
    except Exception as exc:
        if proxy_config and is_retryable_playwright_proxy_error(exc):
            proxy_retry_used = True
            proxy_error_message = str(exc)
            try:
                browser_mode = asyncio.run(
                    _capture_web_screenshot(
                        url=url,
                        output_path=output_path,
                        width=width,
                        height=height,
                        full_page=full_page,
                        wait_ms=wait_ms,
                        browser_executable_path=browser_executable_path,
                        proxy_config=None,
                    )
                )
            except Exception as retry_exc:
                raise RuntimeError(
                    f"网页截图失败：代理模式报错 `{exc}`；直连重试也失败：{retry_exc}"
                ) from retry_exc
        else:
            raise RuntimeError(f"网页截图失败：{exc}") from exc

    file_size = output_path.stat().st_size if output_path.exists() else 0
    media_url = publish_openclaw_media_proxy(output_path, "chinobot-screenshot")
    return {
        "ok": True,
        "message": "✅ 网页截图成功",
        "url": url,
        "file_path": str(output_path),
        "path": str(output_path),
        "media_url": media_url,
        "mediaUrl": media_url,
        "primary_media_url": media_url,
        "primaryMediaUrl": media_url,
        "width": width,
        "height": height,
        "wait_ms": wait_ms,
        "full_page": full_page,
        "file_size_bytes": file_size,
        "browser_executable_path": str(browser_executable_path),
        "browser_mode": browser_mode,
        "transport": {
            "proxy_requested": bool(proxy_config),
            "proxy_server": mask_url_for_debug((proxy_config or {}).get("server", "")),
            "proxy_retry_used": proxy_retry_used,
            "proxy_error": proxy_error_message,
        },
        "note": "截图会由系统自动发送，请不要重复输出链接、文件路径或 [[media:...]]。",
    }


def tool_send_email(params: dict[str, Any]) -> str:
    receiver_email = str(params.get("receiver_email", "")).strip()
    subject = str(params.get("subject", "")).strip()
    content = str(params.get("content", "")).strip()
    attachment_path = params.get("attachment_path")

    sender_email = os.getenv("QQ_EMAIL_SENDER")
    sender_password = os.getenv("QQ_EMAIL_PASSWORD")
    if not sender_email or not sender_password:
        return "❌ 邮件功能未配置（缺少 QQ_EMAIL_SENDER 或 QQ_EMAIL_PASSWORD）"
    if not receiver_email or "@" not in receiver_email:
        return f"❌ 邮箱地址格式不正确: {receiver_email or '[empty]'}"
    if not subject:
        return "❌ 邮件主题不能为空"
    if not content:
        return "❌ 邮件正文不能为空"

    try:
        message = MIMEMultipart()
        message["From"] = sender_email
        message["To"] = receiver_email
        message["Subject"] = Header(subject, "utf-8")
        message.attach(MIMEText(content, "plain", "utf-8"))

        attachment_info = ""
        if attachment_path:
            attachment_file = Path(str(attachment_path)).expanduser().resolve()
            if not attachment_file.exists():
                return f"❌ 附件文件不存在: {attachment_file}"
            with attachment_file.open("rb") as handle:
                attachment = MIMEBase("application", "octet-stream")
                attachment.set_payload(handle.read())
            encoders.encode_base64(attachment)
            attachment.add_header(
                "Content-Disposition",
                f'attachment; filename="{attachment_file.name}"',
            )
            message.attach(attachment)
            attachment_info = f"\n📎 附件：{attachment_file.name} ({attachment_file.stat().st_size / 1024 / 1024:.2f} MB)"

        server = smtplib.SMTP("smtp.qq.com", 587, timeout=10)
        server.ehlo()
        server.starttls()
        server.ehlo()
        server.login(sender_email, sender_password)
        server.sendmail(sender_email, [receiver_email], message.as_string())
        server.quit()

        return (
            "✅ 邮件发送成功！\n"
            f"📧 收件人：{receiver_email}\n"
            f"📝 主题：{subject}{attachment_info}\n"
            f"📤 发件人：{sender_email}"
        )
    except smtplib.SMTPAuthenticationError:
        return "❌ 邮件发送失败：SMTP 认证错误，请检查邮箱配置"
    except smtplib.SMTPException as exc:
        return f"❌ 邮件发送失败：{exc}"
    except Exception as exc:
        return f"❌ 邮件发送失败：{exc}"


def tool_receive_email(params: dict[str, Any]) -> str:
    max_count = int(params.get("max_count", 5) or 5)
    unread_only = bool(params.get("unread_only", True))

    sender_email = os.getenv("QQ_EMAIL_SENDER")
    sender_password = os.getenv("QQ_EMAIL_PASSWORD")
    if not sender_email or not sender_password:
        return "❌ 邮件功能未配置（缺少 QQ_EMAIL_SENDER 或 QQ_EMAIL_PASSWORD）"

    try:
        mail = imaplib.IMAP4_SSL("imap.qq.com", 993)
        mail.login(sender_email, sender_password)
        mail.select("INBOX")
        status, messages = mail.search(None, "UNSEEN" if unread_only else "ALL")
        if status != "OK":
            return "❌ 搜索邮件失败"

        email_ids = messages[0].split()
        if not email_ids:
            return "📭 没有未读邮件" if unread_only else "📭 收件箱为空"

        emails_info = []
        for email_id in email_ids:
            status, msg_data = mail.fetch(email_id, "(RFC822)")
            if status != "OK":
                continue
            msg = email.message_from_bytes(msg_data[0][1])
            subject = decode_mime_header(msg.get("Subject", ""))
            from_header = decode_mime_header(msg.get("From", ""))
            date_header = msg.get("Date", "")
            emails_info.append(
                {
                    "subject": subject or "(无主题)",
                    "from": from_header or "(未知发件人)",
                    "date": date_header,
                }
            )

        mail.logout()
        emails_info = emails_info[-max_count:]
        if not emails_info:
            return "📭 没有可读取的邮件"

        lines = [f"📬 共读取到 {len(emails_info)} 封邮件："]
        for index, info in enumerate(reversed(emails_info), start=1):
            lines.append(
                f"\n{index}. {info['subject']}\n"
                f"   发件人：{info['from']}\n"
                f"   时间：{info['date']}"
            )
        return "\n".join(lines)
    except imaplib.IMAP4.error as exc:
        return f"❌ 读取邮件失败：IMAP 错误：{exc}"
    except Exception as exc:
        return f"❌ 读取邮件失败：{exc}"


def tool_read_word(params: dict[str, Any]) -> str:
    try:
        from docx import Document
    except Exception as exc:
        return f"❌ 读取 Word 失败：缺少 python-docx 依赖（{exc}）"

    file_path = Path(str(params.get("file_path", "")).strip()).expanduser()
    if not file_path.exists():
        return f"❌ 文件不存在: {file_path}"

    try:
        doc = Document(str(file_path))
        paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        if not paragraphs:
            return "❌ Word 文档为空或无法读取"
        summary = (
            "📝 Word 文档信息\n"
            f"📁 文件：{file_path.name}\n"
            f"📊 段落数：{len(paragraphs)}\n"
            f"{'=' * 50}\n\n"
        )
        return summary + "\n\n".join(paragraphs)
    except Exception as exc:
        return f"❌ 读取 Word 失败：{exc}"


def tool_create_word(params: dict[str, Any]) -> str:
    try:
        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    except Exception as exc:
        return f"❌ 创建 Word 失败：缺少 python-docx 依赖（{exc}）"

    file_path = Path(str(params.get("file_path", "")).strip()).name
    title = str(params.get("title", "")).strip()
    content = str(params.get("content", "")).strip()
    if not file_path:
        return "❌ file_path 不能为空"
    if not title:
        return "❌ title 不能为空"
    if not content:
        return "❌ content 不能为空"

    target_path = DOWNLOADS_DIR / file_path
    try:
        doc = Document()
        heading = doc.add_heading(title, level=1)
        heading.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        paragraphs = [part.strip() for part in content.split("\n\n") if part.strip()]
        for para_text in paragraphs:
            doc.add_paragraph(para_text)
        doc.save(str(target_path))
        return (
            "✅ Word 文档创建成功！\n"
            f"📁 文件：{target_path}\n"
            f"📊 大小：{target_path.stat().st_size / 1024:.2f} KB\n"
            f"📝 段落数：{len(paragraphs)}"
        )
    except Exception as exc:
        return f"❌ 创建 Word 失败：{exc}"


def _parse_bilibili(url: str) -> str:
    import httpx

    resolved_url = url
    if "b23.tv" in url or "bili2233.cn" in url:
        with httpx.Client(verify=False, follow_redirects=True, timeout=10) as client:
            resolved_url = str(client.get(url).url)
    if re.fullmatch(r"BV[0-9A-Za-z]{10}", resolved_url):
        resolved_url = f"https://www.bilibili.com/video/{resolved_url}"

    match = re.search(r"video/([^?/ ]+)", resolved_url)
    if not match:
        return f"无法从链接中提取视频ID: {url}"
    video_id = match.group(1)
    api_url = f"https://api.bilibili.com/x/web-interface/view?bvid={video_id}"
    headers = {
        "User-Agent": "Mozilla/5.0",
        "Referer": "https://www.bilibili.com/",
    }
    with httpx.Client(verify=False, timeout=10) as client:
        data = client.get(api_url, headers=headers).json()
    if data.get("code") != 0:
        return f"B站 API 返回错误: {data.get('message', '未知错误')}"
    info = data.get("data") or {}
    stat = info.get("stat") or {}
    lines = [
        "🔗 B站视频解析",
        "━━━━━━━━━━━━━━━━",
        f"📝 标题: {info.get('title', '未知')}",
        f"👤 UP主: {(info.get('owner') or {}).get('name', '未知')}",
    ]
    if stat:
        lines.extend(
            [
                f"👁️ 播放: {stat.get('view', 0):,}",
                f"👍 点赞: {stat.get('like', 0):,}",
                f"💬 评论: {stat.get('reply', 0):,}",
                f"⭐ 收藏: {stat.get('favorite', 0):,}",
            ]
        )
    duration = int(info.get("duration", 0) or 0)
    if duration > 0:
        lines.append(f"⏱️ 时长: {duration // 60}:{duration % 60:02d}")
    desc = str(info.get("desc", "")).strip()
    if desc:
        lines.append(f"📄 简介: {(desc[:150] + '...') if len(desc) > 150 else desc}")
    lines.append(f"🔗 链接: {resolved_url}")
    return "\n".join(lines)


def _parse_github(url: str) -> str:
    import httpx

    match = re.search(r"github\.com/([^/]+)/([^/#?]+)", url)
    if not match:
        return f"无法解析 GitHub 链接: {url}"
    owner, repo = match.group(1), match.group(2).removesuffix(".git")
    api_url = f"https://api.github.com/repos/{owner}/{repo}"
    headers = {"Accept": "application/vnd.github+json", "User-Agent": "openclaw-chinobot-bridge"}
    with httpx.Client(timeout=10) as client:
        response = client.get(api_url, headers=headers)
    if response.status_code != 200:
        return f"GitHub API 请求失败: HTTP {response.status_code}"
    data = response.json()
    lines = [
        "🔗 GitHub 仓库解析",
        "━━━━━━━━━━━━━━━━",
        f"📦 仓库: {data.get('full_name', owner + '/' + repo)}",
        f"📝 描述: {data.get('description') or '无描述'}",
        f"⭐ Stars: {data.get('stargazers_count', 0):,}",
        f"🍴 Forks: {data.get('forks_count', 0):,}",
        f"🐛 Issues: {data.get('open_issues_count', 0):,}",
        f"🌐 主页: {data.get('homepage') or '无'}",
        f"🔗 链接: {url}",
    ]
    return "\n".join(lines)


def tool_parse_link(params: dict[str, Any]) -> str:
    url = str(params.get("url", "")).strip()
    if not url:
        return "❌ url 不能为空"
    try:
        if "bilibili.com" in url or "b23.tv" in url or "bili2233.cn" in url or re.fullmatch(
            r"BV[0-9A-Za-z]{10}",
            url,
        ):
            return _parse_bilibili(url)
        if "github.com" in url:
            return _parse_github(url)
        if "youtube.com" in url or "youtu.be" in url:
            return "YouTube 解析暂未接入当前 bridge（原 chino_bot 依赖 yt_dlp 环境）"
        return f"暂不支持解析此平台的链接: {url}\n\n目前 bridge 已支持：B站、GitHub"
    except Exception as exc:
        return f"解析链接时出错: {exc}\n链接: {url}"


def tool_tavily_search(params: dict[str, Any]) -> dict[str, Any]:
    query = str(params.get("query", "")).strip()
    if not query:
        raise ValueError("query 不能为空")

    search_depth = str(params.get("search_depth", "basic")).strip().lower() or "basic"
    if search_depth not in {"basic", "advanced"}:
        raise ValueError("search_depth 只能是 basic 或 advanced")

    topic = str(params.get("topic", "general")).strip().lower() or "general"
    if topic not in {"general", "news"}:
        raise ValueError("topic 只能是 general 或 news")

    max_results = coerce_int(
        params.get("max_results", params.get("count", 5)),
        default=5,
        minimum=1,
        maximum=10,
    )
    timeout_seconds = coerce_int(params.get("timeout_seconds"), default=20, minimum=5, maximum=60)
    include_answer = coerce_bool(params.get("include_answer"), True)
    include_images = coerce_bool(params.get("include_images"), False)
    include_raw_content = coerce_bool(params.get("include_raw_content"), False)
    dry_run = coerce_bool(params.get("dry_run"), False)

    days_value: int | None = None
    if params.get("days") not in (None, ""):
        days_value = coerce_int(params.get("days"), default=7, minimum=1, maximum=365)

    include_domains = normalize_string_list(params.get("include_domains"))
    exclude_domains = normalize_string_list(params.get("exclude_domains"))
    api_key, api_key_source = resolve_tavily_api_key()
    if not api_key:
        raise RuntimeError("TAVILY_API_KEY 未配置，无法使用 Tavily 搜索")

    proxy_mode = str(params.get("proxy_mode", "auto")).strip().lower() or "auto"
    if proxy_mode not in {"auto", "direct", "project"}:
        raise ValueError("proxy_mode 只能是 auto、direct 或 project")

    request_payload: dict[str, Any] = {
        "query": query,
        "max_results": max_results,
        "search_depth": search_depth,
        "topic": topic,
        "include_answer": include_answer,
        "include_images": include_images,
        "include_raw_content": include_raw_content,
    }
    if topic == "news" and days_value is not None:
        request_payload["days"] = days_value
    if include_domains:
        request_payload["include_domains"] = include_domains
    if exclude_domains:
        request_payload["exclude_domains"] = exclude_domains

    available_proxy = resolve_project_proxy_settings()
    dry_run_payload = {
        "provider": "tavily",
        "query": query,
        "api_key_present": True,
        "api_key_source": api_key_source,
        "project_env_loaded": bool(PROJECT_ENV),
        "proxy_mode": proxy_mode,
        "proxy_available": bool(available_proxy),
        "proxy_http": mask_url_for_debug(available_proxy.get("http", "")),
        "proxy_https": mask_url_for_debug(available_proxy.get("https", "")),
        "use_proxy_env_flag": coerce_bool(os.getenv("USE_PROXY"), False),
        "ssl_verify_default": not coerce_bool(
            os.getenv("TAVILY_DISABLE_SSL_VERIFY"),
            default=True,
        ),
        "request_preview": request_payload,
    }
    if dry_run:
        return dry_run_payload

    attempts: list[bool]
    if proxy_mode == "direct":
        attempts = [False]
    elif proxy_mode == "project":
        attempts = [True]
    else:
        attempts = [False] + ([True] if available_proxy else [])

    errors: list[str] = []
    response_payload: dict[str, Any] | None = None
    transport_meta: dict[str, Any] | None = None
    request_body = {
        **request_payload,
        "api_key": api_key,
    }

    for use_proxy in attempts:
        try:
            response_payload, transport_meta = perform_tavily_request(
                request_payload=request_body,
                timeout_seconds=timeout_seconds,
                use_proxy=use_proxy,
            )
            break
        except Exception as exc:
            mode = "project_proxy" if use_proxy else "direct"
            errors.append(f"{mode}: {exc}")

    if response_payload is None or transport_meta is None:
        raise RuntimeError("；".join(errors) if errors else "Tavily 请求失败")

    answer = response_payload.get("answer")
    if answer is not None and not isinstance(answer, str):
        answer = json.dumps(answer, ensure_ascii=False)

    normalized_results: list[dict[str, Any]] = []
    raw_results = response_payload.get("results")
    if isinstance(raw_results, list):
        for item in raw_results[:max_results]:
            if not isinstance(item, dict):
                continue
            content = maybe_limit_text_output(str(item.get("content", "")).strip(), max_chars=1200)
            result_item: dict[str, Any] = {
                "title": str(item.get("title", "")).strip() or None,
                "url": str(item.get("url", "")).strip() or None,
                "content": content or None,
                "score": item.get("score"),
                "published_date": item.get("published_date") or item.get("publishedDate") or None,
            }
            favicon = str(item.get("favicon", "")).strip()
            if favicon:
                result_item["favicon"] = favicon
            if include_raw_content:
                raw_content = maybe_limit_text_output(
                    str(item.get("raw_content", "")).strip(),
                    max_chars=2000,
                )
                if raw_content:
                    result_item["raw_content"] = raw_content
            normalized_results.append(result_item)

    images: list[str] = []
    if include_images and isinstance(response_payload.get("images"), list):
        for item in response_payload["images"]:
            image_url = str(item).strip()
            if image_url:
                images.append(image_url)

    result: dict[str, Any] = {
        "provider": "tavily",
        "query": query,
        "answer": maybe_limit_text_output(answer.strip(), max_chars=2000) if isinstance(answer, str) else None,
        "results": normalized_results,
        "topic": topic,
        "search_depth": search_depth,
        "max_results": max_results,
        "transport": transport_meta,
        "proxy_mode": proxy_mode,
    }
    if images:
        result["images"] = images[:10]
    if errors:
        result["attempt_errors"] = errors
    return result


TOOLS: dict[str, Callable[[dict[str, Any]], Any]] = {
    "send_email": tool_send_email,
    "receive_email": tool_receive_email,
    "read_pdf": tool_read_pdf,
    "read_word": tool_read_word,
    "create_word": tool_create_word,
    "read_excel": tool_read_excel,
    "create_excel": tool_create_excel,
    "convert_word_to_pdf": tool_convert_word_to_pdf,
    "convert_pdf_to_word": tool_convert_pdf_to_word,
    "parse_link": tool_parse_link,
    "tavily_search": tool_tavily_search,
    "render_html": tool_render_html,
    "web_screenshot": tool_web_screenshot,
}


TOOL_META = {
    "send_email": {
        "bridge_name": "chinobot_send_email",
        "description": "Send email using chino_bot's QQ mail workflow.",
    },
    "receive_email": {
        "bridge_name": "chinobot_receive_email",
        "description": "Read inbox mail using chino_bot's QQ mail workflow.",
    },
    "read_pdf": {
        "bridge_name": "chinobot_read_pdf",
        "description": "Read PDF content using chino_bot's document flow.",
    },
    "read_word": {
        "bridge_name": "chinobot_read_word",
        "description": "Read a Word document using chino_bot's document flow.",
    },
    "create_word": {
        "bridge_name": "chinobot_create_word",
        "description": "Create a Word document using chino_bot's document flow.",
    },
    "read_excel": {
        "bridge_name": "chinobot_read_excel",
        "description": "Read Excel content using chino_bot's document flow.",
    },
    "create_excel": {
        "bridge_name": "chinobot_create_excel",
        "description": "Create an Excel workbook using chino_bot's document flow.",
    },
    "convert_word_to_pdf": {
        "bridge_name": "chinobot_convert_word_to_pdf",
        "description": "Convert a Word document into PDF using chino_bot's document flow.",
    },
    "convert_pdf_to_word": {
        "bridge_name": "chinobot_convert_pdf_to_word",
        "description": "Convert a PDF into Word using chino_bot's document flow.",
    },
    "parse_link": {
        "bridge_name": "chinobot_parse_link",
        "description": "Resolve selected links using chino_bot's link parser.",
    },
    "tavily_search": {
        "bridge_name": "chinobot_tavily_search",
        "description": "Search the live web through chino_bot's Tavily-backed lookup flow.",
    },
    "render_html": {
        "bridge_name": "chinobot_render_html",
        "description": "Render HTML/CSS into a send-ready image using chino_bot's visual workflow.",
    },
    "web_screenshot": {
        "bridge_name": "chinobot_web_screenshot",
        "description": "Capture a webpage screenshot as a send-ready image using chino_bot's visual workflow.",
    },
}


def handle_list_tools() -> dict[str, Any]:
    return ok(
        [
            {
                "tool_name": tool_name,
                **meta,
            }
            for tool_name, meta in TOOL_META.items()
        ]
    )


def handle_invoke(tool_name: str, params: dict[str, Any]) -> dict[str, Any]:
    tool = TOOLS.get(tool_name)
    if tool is None:
        return fail(f"unknown tool: {tool_name}")
    try:
        return ok(tool(params))
    except Exception as exc:
        return fail(str(exc))


def main() -> int:
    parser = argparse.ArgumentParser(description="OpenClaw bridge for selected chino_bot tools")
    subparsers = parser.add_subparsers(dest="command", required=True)

    subparsers.add_parser("list-tools", help="List currently bridged tools")

    invoke_parser = subparsers.add_parser("invoke", help="Invoke one bridged tool")
    invoke_parser.add_argument("--tool", required=True, help="Underlying chino_bot tool name")
    invoke_parser.add_argument("--params-json", default="{}", help="JSON object with tool parameters")

    args = parser.parse_args()

    if args.command == "list-tools":
        payload = handle_list_tools()
    else:
        try:
            params = json.loads(args.params_json or "{}")
            if not isinstance(params, dict):
                payload = fail("params-json must decode to an object")
            else:
                payload = handle_invoke(args.tool, params)
        except json.JSONDecodeError as exc:
            payload = fail(f"invalid params-json: {exc}")

    json.dump(payload, sys.stdout, ensure_ascii=False)
    sys.stdout.write("\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
